#!/usr/bin/env python3
"""
Word Document Figure Migration Script v2
Handles the merged format from pandoc conversion

This script:
1. Identifies all Figure/Table placeholders in the new document
2. Maps them to source content (Mermaid files, screenshots, tables)
3. Generates a comprehensive mapping report
4. Provides instructions for rendering and insertion
"""

import os
import re
from pathlib import Path
from docx import Document

# Paths
BASE_DIR = Path("/Users/scott/Cursor/RWA-Rental-Yield-Tokenisation_v5.1")
ORIGINAL_DOC = BASE_DIR / "DissertationProgress.docx"
NEW_DOC = BASE_DIR / "DissertationProgressFinal.docx"
MERMAID_DIR = BASE_DIR / "Shared_Environment/docs/architecture/diagrams/mermaid"

def find_all_placeholders():
    """Find all Figure and Table placeholders in the new document"""
    print("=== ANALYZING NEW DOCUMENT ===\n")
    
    doc = Document(NEW_DOC)
    
    figures = []
    tables = []
    
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        
        # Match Figure patterns
        fig_match = re.match(r'^Figure\s+([\d\.]+):\s*(.+?)(?:\s*\[|$)', text)
        if fig_match:
            fig_num = fig_match.group(1)
            title = fig_match.group(2).strip()
            
            # Extract mermaid file if present
            mmd_match = re.search(r'([\w-]+\.mmd)', text)
            mmd_file = mmd_match.group(1) if mmd_match else None
            
            # Determine type
            if mmd_file:
                source_type = "Mermaid"
            elif "Screenshot" in text or "screenshot" in text:
                source_type = "Screenshot"
            elif "Chart" in text or "chart" in text:
                source_type = "Chart"
            elif "Diagram" in text or "diagram" in text:
                source_type = "Diagram"
            else:
                source_type = "Unknown"
            
            figures.append({
                'para_index': i,
                'number': fig_num,
                'title': title,
                'mermaid_file': mmd_file,
                'source_type': source_type,
                'full_text': text[:200]
            })
        
        # Match Table patterns
        table_match = re.match(r'^Table\s+([\d\.]+):\s*(.+?)(?:\s*\[|$)', text)
        if table_match:
            tables.append({
                'para_index': i,
                'number': table_match.group(1),
                'title': table_match.group(2).strip(),
                'full_text': text[:200]
            })
    
    return figures, tables

def get_available_mermaid_files():
    """Get list of available Mermaid diagram files"""
    mmd_files = {}
    for mmd in MERMAID_DIR.glob("*.mmd"):
        mmd_files[mmd.name] = mmd
    return mmd_files

def generate_comprehensive_report(figures, tables, mmd_files):
    """Generate a detailed migration report"""
    
    report_path = BASE_DIR / "COMPREHENSIVE_FIGURE_MAPPING.md"
    
    with open(report_path, 'w') as f:
        f.write("# Comprehensive Figure and Table Mapping\n\n")
        f.write(f"Generated: {os.popen('date').read().strip()}\n\n")
        
        # Summary
        f.write("## Executive Summary\n\n")
        f.write(f"| Category | Count |\n")
        f.write(f"|----------|-------|\n")
        f.write(f"| **Total Figures** | {len(figures)} |\n")
        f.write(f"| **Total Tables** | {len(tables)} |\n")
        f.write(f"| **Mermaid Diagrams Available** | {len(mmd_files)} |\n")
        
        # Count by type
        mermaid_count = sum(1 for f in figures if f['source_type'] == 'Mermaid')
        screenshot_count = sum(1 for f in figures if f['source_type'] == 'Screenshot')
        chart_count = sum(1 for f in figures if f['source_type'] == 'Chart')
        diagram_count = sum(1 for f in figures if f['source_type'] == 'Diagram')
        
        f.write(f"| **Mermaid Figures** | {mermaid_count} |\n")
        f.write(f"| **Screenshot Figures** | {screenshot_count} |\n")
        f.write(f"| **Chart Figures** | {chart_count} |\n")
        f.write(f"| **Diagram Figures** | {diagram_count} |\n\n")
        
        # Section 4 Figures
        f.write("---\n\n## Section 4: System Design Figures\n\n")
        sec4_figs = [fig for fig in figures if fig['number'].startswith('4.')]
        f.write(f"**Count: {len(sec4_figs)} figures**\n\n")
        f.write("| Fig # | Title | Source Type | Mermaid File | Para # |\n")
        f.write("|-------|-------|-------------|--------------|--------|\n")
        for fig in sec4_figs:
            mmd = fig['mermaid_file'] or '-'
            f.write(f"| {fig['number']} | {fig['title'][:50]} | {fig['source_type']} | {mmd} | {fig['para_index']} |\n")
        
        # Section 5 Figures
        f.write("\n---\n\n## Section 5: Implementation Figures\n\n")
        sec5_figs = [fig for fig in figures if fig['number'].startswith('5.')]
        f.write(f"**Count: {len(sec5_figs)} figures**\n\n")
        f.write("| Fig # | Title | Source Type | Mermaid File | Para # |\n")
        f.write("|-------|-------|-------------|--------------|--------|\n")
        for fig in sec5_figs:
            mmd = fig['mermaid_file'] or '-'
            f.write(f"| {fig['number']} | {fig['title'][:50]} | {fig['source_type']} | {mmd} | {fig['para_index']} |\n")
        
        # Section 6 Figures
        f.write("\n---\n\n## Section 6: Testing Figures\n\n")
        sec6_figs = [fig for fig in figures if fig['number'].startswith('6.')]
        f.write(f"**Count: {len(sec6_figs)} figures**\n\n")
        f.write("| Fig # | Title | Source Type | Mermaid File | Para # |\n")
        f.write("|-------|-------|-------------|--------------|--------|\n")
        for fig in sec6_figs:
            mmd = fig['mermaid_file'] or '-'
            f.write(f"| {fig['number']} | {fig['title'][:50]} | {fig['source_type']} | {mmd} | {fig['para_index']} |\n")
        
        # Section 7 Figures
        f.write("\n---\n\n## Section 7: Evaluation Figures\n\n")
        sec7_figs = [fig for fig in figures if fig['number'].startswith('7.')]
        f.write(f"**Count: {len(sec7_figs)} figures**\n\n")
        f.write("| Fig # | Title | Source Type | Mermaid File | Para # |\n")
        f.write("|-------|-------|-------------|--------------|--------|\n")
        for fig in sec7_figs:
            mmd = fig['mermaid_file'] or '-'
            f.write(f"| {fig['number']} | {fig['title'][:50]} | {fig['source_type']} | {mmd} | {fig['para_index']} |\n")
        
        # Survey Figures
        f.write("\n---\n\n## Survey Figures (S.x)\n\n")
        survey_figs = [fig for fig in figures if fig['number'].startswith('S.')]
        f.write(f"**Count: {len(survey_figs)} figures**\n\n")
        f.write("| Fig # | Title | Source Type | Para # |\n")
        f.write("|-------|-------|-------------|--------|\n")
        for fig in survey_figs:
            f.write(f"| {fig['number']} | {fig['title'][:50]} | {fig['source_type']} | {fig['para_index']} |\n")
        
        # Interview Figures
        f.write("\n---\n\n## Interview Figures (I.x)\n\n")
        interview_figs = [fig for fig in figures if fig['number'].startswith('I.')]
        f.write(f"**Count: {len(interview_figs)} figures**\n\n")
        f.write("| Fig # | Title | Source Type | Para # |\n")
        f.write("|-------|-------|-------------|--------|\n")
        for fig in interview_figs:
            f.write(f"| {fig['number']} | {fig['title'][:50]} | {fig['source_type']} | {fig['para_index']} |\n")
        
        # Tables
        f.write("\n---\n\n## All Tables\n\n")
        f.write(f"**Count: {len(tables)} tables**\n\n")
        f.write("| Table # | Title | Para # |\n")
        f.write("|---------|-------|--------|\n")
        for tbl in tables:
            f.write(f"| {tbl['number']} | {tbl['title'][:60]} | {tbl['para_index']} |\n")
        
        # Mermaid File Mapping
        f.write("\n---\n\n## Mermaid File Verification\n\n")
        f.write("Checking that all referenced .mmd files exist:\n\n")
        f.write("| Mermaid File | Exists | Used By |\n")
        f.write("|--------------|--------|--------|\n")
        
        referenced_mmd = set()
        for fig in figures:
            if fig['mermaid_file']:
                referenced_mmd.add(fig['mermaid_file'])
        
        for mmd_name in sorted(referenced_mmd):
            exists = "✓" if mmd_name in mmd_files else "✗ MISSING"
            used_by = [f['number'] for f in figures if f['mermaid_file'] == mmd_name]
            f.write(f"| {mmd_name} | {exists} | {', '.join(used_by)} |\n")
        
        # Unused Mermaid files
        f.write("\n### Unused Mermaid Files\n\n")
        unused = set(mmd_files.keys()) - referenced_mmd
        if unused:
            for mmd in sorted(unused):
                f.write(f"- {mmd}\n")
        else:
            f.write("All Mermaid files are referenced.\n")
        
        # Action Plan
        f.write("\n---\n\n## ACTION PLAN\n\n")
        f.write("### Step 1: Render Mermaid Diagrams\n\n")
        f.write("```bash\n")
        f.write("# Install mermaid-cli if not already installed\n")
        f.write("npm install -g @mermaid-js/mermaid-cli\n\n")
        f.write("# Render all Mermaid diagrams\n")
        f.write("cd /Users/scott/Cursor/RWA-Rental-Yield-Tokenisation_v5.1\n")
        f.write("mkdir -p rendered_diagrams\n\n")
        for mmd_name in sorted(referenced_mmd):
            base_name = mmd_name.replace('.mmd', '')
            f.write(f"mmdc -i Shared_Environment/docs/architecture/diagrams/mermaid/{mmd_name} -o rendered_diagrams/{base_name}.png -w 2000\n")
        f.write("```\n\n")
        
        f.write("### Step 2: Extract Screenshots from Original Document\n\n")
        f.write("The 10 images have been extracted to `extracted_images/` folder.\n")
        f.write("Map them to the screenshot placeholders manually.\n\n")
        
        f.write("### Step 3: Create Charts\n\n")
        f.write("Charts need to be created/extracted from original document or created fresh:\n\n")
        for fig in figures:
            if fig['source_type'] == 'Chart':
                f.write(f"- Figure {fig['number']}: {fig['title']}\n")
        
        f.write("\n### Step 4: Insert into Word Document\n\n")
        f.write("Use the paragraph indices to locate each placeholder and insert the rendered image.\n\n")
    
    print(f"✓ Report saved to: {report_path}")
    return report_path

def main():
    print("=" * 70)
    print("WORD DOCUMENT FIGURE MIGRATION - COMPREHENSIVE ANALYSIS")
    print("=" * 70 + "\n")
    
    # Find all placeholders
    figures, tables = find_all_placeholders()
    
    print(f"Found {len(figures)} figures and {len(tables)} tables\n")
    
    # Get available Mermaid files
    mmd_files = get_available_mermaid_files()
    print(f"Found {len(mmd_files)} Mermaid diagram files\n")
    
    # Generate report
    report_path = generate_comprehensive_report(figures, tables, mmd_files)
    
    # Print summary
    print("\n" + "=" * 70)
    print("SUMMARY BY SECTION")
    print("=" * 70)
    
    sections = {
        '4': [f for f in figures if f['number'].startswith('4.')],
        '5': [f for f in figures if f['number'].startswith('5.')],
        '6': [f for f in figures if f['number'].startswith('6.')],
        '7': [f for f in figures if f['number'].startswith('7.')],
        'S': [f for f in figures if f['number'].startswith('S.')],
        'I': [f for f in figures if f['number'].startswith('I.')]
    }
    
    for sec, figs in sections.items():
        if figs:
            mermaid = sum(1 for f in figs if f['source_type'] == 'Mermaid')
            screenshot = sum(1 for f in figs if f['source_type'] == 'Screenshot')
            chart = sum(1 for f in figs if f['source_type'] == 'Chart')
            other = len(figs) - mermaid - screenshot - chart
            print(f"\nSection {sec}: {len(figs)} figures")
            print(f"  - Mermaid: {mermaid}")
            print(f"  - Screenshots: {screenshot}")
            print(f"  - Charts: {chart}")
            print(f"  - Other: {other}")
    
    print(f"\nTables: {len(tables)}")
    
    print(f"\n✓ See detailed report: {report_path}")

if __name__ == "__main__":
    main()





